import os
import uuid
import sqlite3
import re
import io
from datetime import datetime
from pathlib import Path

from flask import Flask, request, jsonify, send_from_directory, render_template_string
from PIL import Image

try:
    from docx import Document as DocxDocument
    DOCX_SUPPORT = True
except ImportError:
    DOCX_SUPPORT = False

app = Flask(__name__)

BASE_DIR   = Path(__file__).parent
DB_PATH    = BASE_DIR / "logboek.db"
FOTOS_DIR  = BASE_DIR / "fotos"
FOTOS_DIR.mkdir(exist_ok=True)

MAX_IMAGE_WIDTH = 1600
JPEG_QUALITY    = 82


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with get_db() as conn:
        conn.executescript("""
            CREATE TABLE IF NOT EXISTS instellingen (
                sleutel TEXT PRIMARY KEY,
                waarde  TEXT
            );
            CREATE TABLE IF NOT EXISTS document (
                id          INTEGER PRIMARY KEY CHECK (id = 1),
                html        TEXT NOT NULL DEFAULT '',
                bijgewerkt  TEXT NOT NULL DEFAULT ''
            );
            INSERT OR IGNORE INTO document(id, html, bijgewerkt) VALUES(1,'','');
        """)


init_db()


def compress_image(data: bytes) -> bytes:
    img = Image.open(io.BytesIO(data))
    if img.mode not in ('RGB', 'L'):
        img = img.convert('RGB')
    if img.width > MAX_IMAGE_WIDTH:
        ratio = MAX_IMAGE_WIDTH / img.width
        img = img.resize((MAX_IMAGE_WIDTH, int(img.height * ratio)), Image.LANCZOS)
    out = io.BytesIO()
    img.save(out, format='JPEG', quality=JPEG_QUALITY, optimize=True)
    return out.getvalue()


def html_to_plain(html: str) -> str:
    text = re.sub(r'<img[^>]*>', ' [afbeelding] ', html)
    text = re.sub(r'<[^>]+>', ' ', text)
    for ent, ch in [('&nbsp;',' '),('&lt;','<'),('&gt;','>'),('&amp;','&')]:
        text = text.replace(ent, ch)
    return re.sub(r'\s+', ' ', text).strip()


# ── API ──────────────────────────────────────────────────────────────────────

@app.route('/')
def index():
    return render_template_string(HTML)


@app.route('/api/instellingen', methods=['GET'])
def get_instellingen():
    with get_db() as conn:
        rows = conn.execute("SELECT sleutel, waarde FROM instellingen").fetchall()
    return jsonify({r['sleutel']: r['waarde'] for r in rows})


@app.route('/api/instellingen', methods=['POST'])
def set_instellingen():
    data = request.get_json(force=True)
    with get_db() as conn:
        for k, v in data.items():
            conn.execute(
                "INSERT INTO instellingen(sleutel,waarde) VALUES(?,?) "
                "ON CONFLICT(sleutel) DO UPDATE SET waarde=excluded.waarde",
                (k, v))
    return jsonify(ok=True)


@app.route('/api/document', methods=['GET'])
def get_document():
    with get_db() as conn:
        row = conn.execute("SELECT html, bijgewerkt FROM document WHERE id=1").fetchone()
    return jsonify(html=row['html'], bijgewerkt=row['bijgewerkt'])


@app.route('/api/document', methods=['PUT'])
def put_document():
    data = request.get_json(force=True)
    html = data.get('html', '')
    nu   = datetime.now().isoformat(timespec='seconds')
    with get_db() as conn:
        conn.execute("UPDATE document SET html=?, bijgewerkt=? WHERE id=1", (html, nu))
    return jsonify(ok=True, bijgewerkt=nu)


@app.route('/api/zoek')
def zoek():
    q = request.args.get('q', '').strip()
    if not q:
        return jsonify(resultaten=[], totaal=0)
    with get_db() as conn:
        row = conn.execute("SELECT html FROM document WHERE id=1").fetchone()
    html  = row['html'] if row else ''
    plain = html_to_plain(html)
    aantal = plain.lower().count(q.lower())
    return jsonify(resultaten=aantal, totaal=len(plain))


@app.route('/api/foto', methods=['POST'])
def upload_foto():
    if 'file' not in request.files:
        return jsonify(error='geen bestand'), 400
    f    = request.files['file']
    data = compress_image(f.read())
    naam = f"{uuid.uuid4().hex}.jpg"
    (FOTOS_DIR / naam).write_bytes(data)
    return jsonify(url=f'/fotos/{naam}')


@app.route('/fotos/<path:naam>')
def serve_foto(naam):
    return send_from_directory(FOTOS_DIR, naam)


@app.route('/api/open', methods=['POST'])
def open_bestand():
    if 'file' not in request.files:
        return jsonify(error='geen bestand'), 400
    f   = request.files['file']
    ext = Path(f.filename or '').suffix.lower()

    if ext == '.txt':
        tekst = f.read().decode('utf-8', errors='replace')
        stijl = "font-size:12pt;font-family:'Segoe UI',Arial,sans-serif;"
        delen = []
        for regel in tekst.splitlines():
            esc = regel.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')
            delen.append(f'<p style="{stijl}">{esc or "<br>"}</p>')
        return jsonify(html=''.join(delen))

    if ext == '.docx':
        if not DOCX_SUPPORT:
            return jsonify(error='python-docx niet geïnstalleerd'), 501
        doc      = DocxDocument(io.BytesIO(f.read()))
        basestijl = "font-size:12pt;font-family:'Segoe UI',Arial,sans-serif;"
        delen = []
        for para in doc.paragraphs:
            pstijl = (para.style.name or '').lower()
            tag    = 'h1' if 'heading 1' in pstijl else 'h2' if 'heading 2' in pstijl else 'p'
            inhoud = ''
            for run in para.runs:
                t = run.text.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')
                if run.bold:      t = f'<strong>{t}</strong>'
                if run.italic:    t = f'<em>{t}</em>'
                if run.underline: t = f'<u>{t}</u>'
                inhoud += t
            css = '' if tag != 'p' else f' style="{basestijl}"'
            delen.append(f'<{tag}{css}>{inhoud or "<br>"}</{tag}>')
        return jsonify(html=''.join(delen))

    return jsonify(error=f'Bestandstype {ext} niet ondersteund'), 415


@app.route('/api/export')
def export_doc():
    formaat = request.args.get('formaat', 'json')
    with get_db() as conn:
        row = conn.execute("SELECT html, bijgewerkt FROM document WHERE id=1").fetchone()
    html  = row['html'] if row else ''
    plain = html_to_plain(html)

    if formaat == 'txt':
        from flask import Response
        return Response(plain, mimetype='text/plain; charset=utf-8')
    if formaat == 'html':
        from flask import Response
        volledige = f"""<!DOCTYPE html>
<html lang="nl"><head><meta charset="utf-8">
<title>Logboek export</title>
<style>body{{font-family:'Segoe UI',Arial,sans-serif;max-width:860px;margin:40px auto;padding:0 20px}}
img{{max-width:100%}}.logvak{{border-top:2px solid #555;border-bottom:2px solid #555;padding:10px 0;margin:18px 0;font-family:'Courier New',monospace}}</style>
</head><body>{html}</body></html>"""
        return Response(volledige, mimetype='text/html; charset=utf-8')
    return jsonify(
        export=datetime.now().isoformat(timespec='seconds'),
        bijgewerkt=row['bijgewerkt'],
        platte_tekst=plain,
        html=html
    )


# ── HTML/CSS/JS ──────────────────────────────────────────────────────────────

HTML = r"""<!DOCTYPE html>
<html lang="nl">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Logboek</title>
<style>
*, *::before, *::after { box-sizing: border-box; margin: 0; padding: 0; }

body {
  font-family: 'Segoe UI', Arial, sans-serif;
  background: #ababab;
  min-height: 100vh;
  display: flex;
  flex-direction: column;
}

/* ── menubalk (WordPad-stijl) ── */
#menubar {
  background: #f0f0f0;
  border-bottom: 1px solid #999;
  padding: 2px 4px;
  font-size: 12px;
  display: flex;
  gap: 2px;
}
#menubar span {
  padding: 3px 8px;
  border-radius: 2px;
  cursor: default;
  color: #222;
}
#menubar span:hover { background: #d0d8e8; }

/* ── toolbar ── */
#toolbar {
  background: linear-gradient(to bottom, #fefefe 0%, #e8e8e8 100%);
  border-bottom: 1px solid #aaa;
  padding: 4px 8px;
  display: flex;
  flex-wrap: wrap;
  gap: 3px;
  align-items: center;
  box-shadow: 0 1px 2px rgba(0,0,0,.12);
}

.tb-groep {
  display: flex;
  gap: 2px;
  align-items: center;
  padding: 0 4px;
  border-right: 1px solid #ccc;
}
.tb-groep:last-child { border-right: none; margin-left: auto; }

#toolbar button {
  height: 26px;
  min-width: 26px;
  border: 1px solid transparent;
  border-radius: 2px;
  background: transparent;
  cursor: pointer;
  font-size: 13px;
  padding: 0 5px;
  color: #222;
}
#toolbar button:hover  { background: #d4e0f5; border-color: #90b0e0; }
#toolbar button:active { background: #b8ccf0; }

#toolbar select {
  height: 24px;
  border: 1px solid #bbb;
  border-radius: 2px;
  background: white;
  font-size: 12px;
  padding: 0 2px;
}
#toolbar input[type=color] {
  width: 26px; height: 26px;
  border: 1px solid #bbb;
  border-radius: 2px;
  padding: 1px;
  cursor: pointer;
}

/* NIEUW VAK knop — opvallend */
#btn-nieuwvak {
  background: #2255bb !important;
  color: white !important;
  border-color: #1a3f99 !important;
  font-weight: 600;
  padding: 0 12px !important;
  height: 26px;
  border-radius: 2px;
}
#btn-nieuwvak:hover { background: #1a3f99 !important; }

#zoek-input {
  height: 24px;
  border: 1px solid #bbb;
  border-radius: 2px;
  padding: 0 8px;
  font-size: 12px;
  width: 160px;
}
#status-balk {
  font-size: 11px;
  color: #666;
  padding: 0 6px;
  white-space: nowrap;
}

/* ── papier ── */
#papier-wrap {
  flex: 1;
  overflow-y: auto;
  padding: 24px 0 60px;
  display: flex;
  flex-direction: column;
  align-items: center;
}

#document {
  width: 794px;          /* A4-breedte bij 96dpi */
  min-height: 1123px;
  background: white;
  box-shadow: 0 2px 8px rgba(0,0,0,.35);
  padding: 60px 72px;
  outline: none;
  font-size: 13pt;
  line-height: 1.6;
  color: #111;
  caret-color: #000;
  white-space: pre-wrap;
  word-wrap: break-word;
}

/* vak-blok */
.logvak {
  margin: 18px 0;
  border-top: 2px solid #555;
  border-bottom: 2px solid #555;
  padding: 10px 0;
  font-family: 'Courier New', monospace;
  font-size: 11pt;
}
.logvak .vak-scheiding {
  color: #555;
  letter-spacing: 1px;
  user-select: none;
  font-size: 10pt;
  line-height: 1.2;
}
.logvak .vak-rij {
  display: flex;
  gap: 12px;
  padding: 3px 0;
}
.logvak .vak-label {
  color: #777;
  min-width: 100px;
  font-size: 10pt;
  user-select: none;
}
.logvak .vak-inhoud {
  flex: 1;
  outline: none;
  min-height: 1.4em;
  font-family: inherit;
  font-size: inherit;
  color: #111;
  border-bottom: 1px dashed #ccc;
}
.logvak .vak-datum-tijd { color: #333; font-weight: 600; }

#document img {
  max-width: 100%;
  height: auto;
  display: inline-block;
  margin: 6px 0;
  cursor: pointer;
  vertical-align: bottom;
}
#document img.geselecteerd {
  outline: 2px solid #2255bb;
}

/* resize-wrapper */
.img-resize-wrap {
  display: inline-block;
  position: relative;
  line-height: 0;
  margin: 6px 0;
}
.img-resize-wrap img {
  display: block;
  margin: 0;
}
.resize-greep {
  position: absolute;
  width: 10px; height: 10px;
  background: #2255bb;
  border: 1px solid white;
  border-radius: 2px;
  z-index: 10;
  cursor: se-resize;
}
.resize-greep.rb { bottom: -5px; right: -5px; cursor: se-resize; }
.resize-greep.lb { bottom: -5px; left:  -5px; cursor: sw-resize; }
.resize-greep.rt { top:    -5px; right: -5px; cursor: ne-resize; }
.resize-greep.lt { top:    -5px; left:  -5px; cursor: nw-resize; }

/* ── naam-modal ── */
#naam-modal {
  display: none;
  position: fixed; inset: 0;
  background: rgba(0,0,0,.5);
  z-index: 300;
  align-items: center;
  justify-content: center;
}
#naam-modal.zichtbaar { display: flex; }
.modal-doos {
  background: white;
  border-radius: 4px;
  padding: 28px 36px;
  max-width: 340px;
  width: 90%;
  box-shadow: 0 6px 24px rgba(0,0,0,.3);
}
.modal-doos h2 { font-size: 16px; margin-bottom: 8px; }
.modal-doos p  { font-size: 12px; color: #555; margin-bottom: 16px; }
.modal-doos input {
  width: 100%; padding: 7px 10px;
  border: 1px solid #bbb; border-radius: 3px;
  font-size: 14px; margin-bottom: 14px;
}
.modal-doos button {
  width: 100%; padding: 8px;
  background: #2255bb; color: white;
  border: none; border-radius: 3px;
  font-size: 14px; cursor: pointer;
}
.modal-doos button:hover { background: #1a3f99; }

/* ── vak verwijderen & inklappen ── */
.logvak {
  position: relative;
}
.vak-sluit {
  display: none;
  position: absolute;
  top: 2px; right: 2px;
  width: 20px; height: 20px;
  background: #c0392b;
  color: white;
  border: none;
  border-radius: 3px;
  font-size: 13px;
  line-height: 1;
  cursor: pointer;
  z-index: 5;
}
.logvak:hover .vak-sluit { display: block; }
.vak-sluit:hover { background: #922b21; }

.vak-scheiding-top {
  cursor: pointer;
  user-select: none;
}
.vak-scheiding-top:hover { color: #2255bb; }
.logvak.ingeklapt .vak-rij { display: none; }
.logvak.ingeklapt .vak-scheiding-top::after {
  content: ' ▶ ingeklapt';
  font-size: 9pt;
  color: #999;
}
.logvak:not(.ingeklapt) .vak-scheiding-top::after {
  content: ' ▼';
  font-size: 9pt;
  color: #aaa;
}

/* ── zoek & vervang paneel ── */
#vervang-paneel {
  display: none;
  position: fixed;
  top: 80px; right: 24px;
  background: white;
  border: 1px solid #bbb;
  border-radius: 4px;
  box-shadow: 0 4px 16px rgba(0,0,0,.2);
  padding: 14px 16px;
  z-index: 250;
  min-width: 280px;
  font-size: 13px;
}
#vervang-paneel.zichtbaar { display: block; }
#vervang-paneel h3 { font-size: 13px; margin-bottom: 10px; color: #333; }
#vervang-paneel input {
  width: 100%; padding: 5px 8px;
  border: 1px solid #bbb; border-radius: 3px;
  font-size: 13px; margin-bottom: 6px;
}
#vervang-paneel .pnl-knoppen {
  display: flex; gap: 6px; margin-top: 8px;
}
#vervang-paneel button {
  flex: 1; padding: 5px 8px;
  border: 1px solid #bbb; border-radius: 3px;
  background: #f0f0f0; cursor: pointer; font-size: 12px;
}
#vervang-paneel button:hover { background: #d4e0f5; }
#vervang-paneel .sluit-pnl {
  position: absolute; top: 6px; right: 8px;
  background: none; border: none; cursor: pointer;
  font-size: 16px; color: #888; flex: none; padding: 0;
}

/* ── print ── */
@media print {
  #menubar, #toolbar, #vervang-paneel { display: none !important; }
  body { background: white !important; }
  #papier-wrap { padding: 0 !important; }
  #document { box-shadow: none !important; width: 100% !important; }
}

/* drag-over */
body.drag-over::after {
  content: 'Laat los om afbeelding in te voegen';
  position: fixed; inset: 0;
  background: rgba(34,85,187,.15);
  border: 4px dashed #2255bb;
  z-index: 200;
  display: flex; align-items: center; justify-content: center;
  font-size: 20px; color: #2255bb; pointer-events: none;
}

@media (max-width: 860px) {
  #document { width: 100%; padding: 32px 20px; }
}
</style>
</head>
<body>

<!-- naam-modal -->
<div id="naam-modal">
  <div class="modal-doos">
    <h2>Welkom bij Logboek</h2>
    <p>Vul je naam in. Deze wordt automatisch ingevuld bij elk nieuw vak.</p>
    <input id="naam-input" type="text" placeholder="Jouw naam" autocomplete="name">
    <button onclick="slaaNaamOpEnUpdate()">Opslaan &amp; beginnen</button>
  </div>
</div>

<!-- naam-modal wordt ook hergebruikt voor naam wijzigen -->

<!-- zoek & vervang paneel -->
<div id="vervang-paneel">
  <button class="sluit-pnl" onclick="sluitVervang()">&#10005;</button>
  <h3>Zoeken &amp; vervangen</h3>
  <input id="vervang-zoek"    type="text" placeholder="Zoeken…">
  <input id="vervang-door"    type="text" placeholder="Vervangen door…">
  <div class="pnl-knoppen">
    <button onclick="vervangEen()">Vervang</button>
    <button onclick="vervangAlles()">Alles vervangen</button>
  </div>
</div>

<!-- menubalk -->
<div id="menubar">
  <span onclick="kiesBestand()">&#128194; Openen (.txt / .docx)</span>
  <span onclick="exporteerJSON()">&#128229; Export JSON</span>
  <span onclick="exporteerHTML()">&#128196; Export HTML</span>
  <span onclick="exporteerTXT()">&#128203; Export TXT</span>
  <span onclick="window.print()">&#128438; Afdrukken</span>
  <span onclick="wisselNaam()" id="naam-menu-label" style="margin-left:auto;color:#555;">&#128100; ...</span>
  <input type="file" id="bestand-input" accept=".txt,.docx" style="display:none" onchange="openBestand(this)">
</div>

<!-- toolbar -->
<div id="toolbar">

  <!-- Nieuw vak -->
  <div class="tb-groep">
    <button id="btn-nieuwvak" onclick="voegVakIn()" title="Nieuw logvak invoegen met datum, tijd en naam">
      &#43; Nieuw vak
    </button>
  </div>

  <!-- opmaak -->
  <div class="tb-groep">
    <button onclick="fmt('bold')"      title="Vet (Ctrl+B)"><b>V</b></button>
    <button onclick="fmt('italic')"    title="Cursief (Ctrl+I)"><i>C</i></button>
    <button onclick="fmt('underline')" title="Onderstrepen (Ctrl+U)"><u>O</u></button>
  </div>

  <!-- lettertype -->
  <div class="tb-groep">
    <select id="sel-font" title="Lettertype" onchange="fmtFont(this.value)">
      <option value="'Segoe UI',sans-serif">Segoe UI</option>
      <option value="Arial,sans-serif">Arial</option>
      <option value="'Times New Roman',serif">Times New Roman</option>
      <option value="'Courier New',monospace">Courier New</option>
      <option value="Georgia,serif">Georgia</option>
      <option value="Verdana,sans-serif">Verdana</option>
    </select>
    <select id="sel-size" title="Lettergrootte" onchange="fmtSize(this.value)">
      <option value="9">9</option>
      <option value="10">10</option>
      <option value="11">11</option>
      <option value="12" selected>12</option>
      <option value="14">14</option>
      <option value="16">16</option>
      <option value="18">18</option>
      <option value="20">20</option>
      <option value="24">24</option>
    </select>
    <input type="color" id="kleur-kiezer" value="#000000" title="Tekstkleur"
           onchange="fmtKleur(this.value)">
  </div>

  <!-- lijsten + uitlijning -->
  <div class="tb-groep">
    <button onclick="fmt('insertUnorderedList')" title="Opsomming">&#8226;</button>
    <button onclick="fmt('insertOrderedList')"   title="Genummerd">1.</button>
    <button onclick="fmt('justifyLeft')"   title="Links">&#8676;</button>
    <button onclick="fmt('justifyCenter')" title="Midden">&#8596;</button>
    <button onclick="fmt('justifyRight')"  title="Rechts">&#8677;</button>
  </div>

  <!-- foto -->
  <div class="tb-groep">
    <button onclick="kiesFoto()" title="Afbeelding invoegen">&#128247; Foto</button>
    <input type="file" id="foto-input" accept="image/*" style="display:none" onchange="uploadFoto(this)">
  </div>

  <!-- zoek + navigatie + status -->
  <div class="tb-groep">
    <input id="zoek-input" type="search" placeholder="Zoeken… (Ctrl+H vervangen)" oninput="zoek(this.value)" style="width:190px">
    <button onclick="zoekNavigeer(-1)" title="Vorige">&#8249;</button>
    <button onclick="zoekNavigeer(1)"  title="Volgende">&#8250;</button>
    <span id="zoek-teller" style="font-size:11px;color:#666;padding:0 4px;white-space:nowrap;"></span>
    <span id="status-balk">klaar</span>
  </div>

</div>

<!-- papier -->
<div id="papier-wrap">
  <div id="document" contenteditable="true"
       spellcheck="true"
       onpaste="verwerkPlak(event)"
       oninput="documentGewijzigd()">
  </div>
</div>

<script>
'use strict';

let gebruikersnaam = '';
let bewaarTimer    = null;
let zoekMarkeringen = [];

const doc = document.getElementById('document');

// ── init ─────────────────────────────────────────────────────────────────────

async function init() {
  const inst = await api('GET', '/api/instellingen');
  if (inst.naam) {
    gebruikersnaam = inst.naam;
  } else {
    document.getElementById('naam-modal').classList.add('zichtbaar');
    return;
  }
  await laadDocument();
}

async function slaaNaamOp() {
  const v = document.getElementById('naam-input').value.trim();
  if (!v) return;
  gebruikersnaam = v;
  await api('POST', '/api/instellingen', { naam: v });
  document.getElementById('naam-modal').classList.remove('zichtbaar');
  await laadDocument();
}
document.getElementById('naam-input').addEventListener('keydown', e => {
  if (e.key === 'Enter') slaaNaamOp();
});

async function laadDocument() {
  const data = await api('GET', '/api/document');
  if (data.html) {
    doc.innerHTML = data.html;
  } else {
    doc.innerHTML = '<p><br></p>';
  }
  setStatus(data.bijgewerkt ? 'Opgeslagen: ' + netjesTijd(data.bijgewerkt) : 'Nieuw document');
  document.getElementById('naam-menu-label').textContent = '👤 ' + gebruikersnaam;
}

async function slaaNaamOpEnUpdate() {
  await slaaNaamOp();
  document.getElementById('naam-menu-label').textContent = '👤 ' + gebruikersnaam;
}

// ── nieuw vak invoegen ────────────────────────────────────────────────────────

function voegVakIn() {
  const nu       = new Date();
  const datum    = nu.toLocaleDateString('nl-NL', { weekday:'long', year:'numeric', month:'long', day:'numeric' });
  const tijd     = nu.toLocaleTimeString('nl-NL', { hour:'2-digit', minute:'2-digit' });
  const naam     = escHtml(gebruikersnaam);

  // bouw het vak als HTML
  const vakHtml = `
<div class="logvak" contenteditable="false">
  <button class="vak-sluit" title="Vak verwijderen" onclick="verwijderVak(this)">&#10005;</button>
  <div class="vak-scheiding vak-scheiding-top" onclick="klapVak(this)" title="Klik om in/uit te klappen">════════════════════════════════════════════════════════════════</div>
  <div class="vak-rij">
    <span class="vak-label">Datum/Tijd</span>
    <span class="vak-datum-tijd vak-inhoud">${datum} — ${tijd}</span>
  </div>
  <div class="vak-rij">
    <span class="vak-label">Informatie</span>
    <span class="vak-inhoud" contenteditable="true" data-placeholder="Typ hier je informatie…"></span>
  </div>
  <div class="vak-rij">
    <span class="vak-label">Naam</span>
    <span class="vak-inhoud" contenteditable="true">${naam}</span>
  </div>
  <div class="vak-scheiding">════════════════════════════════════════════════════════════════</div>
</div><p><br></p>`;

  // voeg in op cursorpositie of aan einde
  doc.focus();
  const sel = window.getSelection();
  if (sel && sel.rangeCount) {
    const range = sel.getRangeAt(0);
    // zorg dat we binnen #document zijn
    if (doc.contains(range.commonAncestorContainer)) {
      range.deleteContents();
      const tmp = document.createElement('div');
      tmp.innerHTML = vakHtml;
      const frag = document.createDocumentFragment();
      let firstEl = null;
      while (tmp.firstChild) {
        if (!firstEl) firstEl = tmp.firstChild;
        frag.appendChild(tmp.firstChild);
      }
      range.insertNode(frag);
      // focus op informatie-veld
      if (firstEl) {
        const infoVeld = firstEl.querySelector('.vak-inhoud[contenteditable=true]');
        if (infoVeld) {
          infoVeld.focus();
          const r = document.createRange();
          r.selectNodeContents(infoVeld);
          r.collapse(false);
          sel.removeAllRanges();
          sel.addRange(r);
        }
      }
      documentGewijzigd();
      return;
    }
  }
  // fallback: voeg toe aan einde
  doc.insertAdjacentHTML('beforeend', vakHtml);
  documentGewijzigd();
  const vakken = doc.querySelectorAll('.logvak');
  const laatste = vakken[vakken.length - 1];
  if (laatste) {
    const infoVeld = laatste.querySelector('.vak-inhoud[contenteditable=true]');
    if (infoVeld) infoVeld.focus();
    laatste.scrollIntoView({ behavior: 'smooth', block: 'center' });
  }
}

// ── bewaren ───────────────────────────────────────────────────────────────────

function documentGewijzigd() {
  setStatus('Niet opgeslagen…');
  clearTimeout(bewaarTimer);
  bewaarTimer = setTimeout(bewaar, 1500);
}

async function bewaar() {
  clearTimeout(bewaarTimer);
  setStatus('Bezig met opslaan…');
  const res = await api('PUT', '/api/document', { html: doc.innerHTML });
  setStatus('Opgeslagen: ' + netjesTijd(res.bijgewerkt));
}

// ── opmaak ────────────────────────────────────────────────────────────────────

function fmt(cmd)        { document.execCommand(cmd, false, null); }
function fmtFont(f)      { document.execCommand('fontName', false, f); }
function fmtKleur(k)     { document.execCommand('foreColor', false, k); }
function fmtSize(pt) {
  document.execCommand('fontSize', false, '7');
  doc.querySelectorAll('font[size="7"]').forEach(s => {
    s.removeAttribute('size');
    s.style.fontSize = pt + 'pt';
  });
}

// ── foto ──────────────────────────────────────────────────────────────────────

function kiesFoto() { document.getElementById('foto-input').click(); }

async function uploadFoto(input) {
  const file = input.files[0];
  if (!file) return;
  input.value = '';
  await inserteerFotoBestand(file);
}

async function inserteerFotoBestand(file) {
  const form = new FormData();
  form.append('file', file);
  const res  = await fetch('/api/foto', { method:'POST', body: form });
  const data = await res.json();
  if (data.url) {
    doc.focus();
    document.execCommand('insertImage', false, data.url);
    documentGewijzigd();
  }
}

async function verwerkPlak(e) {
  const items = (e.clipboardData || window.clipboardData).items;
  for (const item of items) {
    if (item.type.startsWith('image/')) {
      e.preventDefault();
      await inserteerFotoBestand(item.getAsFile());
      return;
    }
  }
}

// drag & drop
document.addEventListener('dragover', e => {
  if ([...e.dataTransfer.items].some(i => i.kind==='file' && i.type.startsWith('image/'))) {
    e.preventDefault();
    document.body.classList.add('drag-over');
  }
});
document.addEventListener('dragleave', () => document.body.classList.remove('drag-over'));
document.addEventListener('drop', async e => {
  document.body.classList.remove('drag-over');
  for (const file of e.dataTransfer.files) {
    if (file.type.startsWith('image/')) {
      e.preventDefault();
      await inserteerFotoBestand(file);
    }
  }
});

// ── bestand openen ────────────────────────────────────────────────────────────

function kiesBestand() { document.getElementById('bestand-input').click(); }

async function openBestand(input) {
  const file = input.files[0];
  if (!file) return;
  input.value = '';
  const form = new FormData();
  form.append('file', file);
  const res  = await fetch('/api/open', { method:'POST', body: form });
  if (!res.ok) { alert('Fout bij openen'); return; }
  const data = await res.json();
  // voeg in op cursorpositie
  doc.focus();
  document.execCommand('insertHTML', false, data.html);
  documentGewijzigd();
}

// ── vak verwijderen & inklappen ──────────────────────────────────────────────

function verwijderVak(knop) {
  const vak = knop.closest('.logvak');
  if (!vak) return;
  if (!confirm('Dit logvak verwijderen?')) return;
  vak.remove();
  documentGewijzigd();
}

function klapVak(scheiding) {
  const vak = scheiding.closest('.logvak');
  if (vak) {
    vak.classList.toggle('ingeklapt');
    documentGewijzigd();
  }
}

// ── naam wisselen ─────────────────────────────────────────────────────────────

function wisselNaam() {
  const input = document.getElementById('naam-input');
  input.value = gebruikersnaam;
  document.getElementById('naam-modal').classList.add('zichtbaar');
  input.focus();
  input.select();
}

// ── zoeken ────────────────────────────────────────────────────────────────────

let zoekIndex = -1;

function zoek(q) {
  doc.querySelectorAll('mark.zoek-mark').forEach(m => m.replaceWith(...m.childNodes));
  doc.normalize();
  zoekIndex = -1;
  document.getElementById('zoek-teller').textContent = '';

  if (!q || q.length < 2) return;

  const walker = document.createTreeWalker(doc, NodeFilter.SHOW_TEXT);
  const nodes  = [];
  let node;
  while ((node = walker.nextNode())) nodes.push(node);

  const re = new RegExp(escRegex(q), 'gi');
  nodes.forEach(n => {
    if (!n.nodeValue.match(re)) return;
    const frag = document.createDocumentFragment();
    let last = 0, m;
    re.lastIndex = 0;
    while ((m = re.exec(n.nodeValue)) !== null) {
      frag.appendChild(document.createTextNode(n.nodeValue.slice(last, m.index)));
      const mark = document.createElement('mark');
      mark.className = 'zoek-mark';
      mark.style.background = '#fff176';
      mark.textContent = m[0];
      frag.appendChild(mark);
      last = m.index + m[0].length;
    }
    frag.appendChild(document.createTextNode(n.nodeValue.slice(last)));
    n.parentNode.replaceChild(frag, n);
  });

  const alle = doc.querySelectorAll('mark.zoek-mark');
  if (alle.length) {
    zoekIndex = 0;
    markeerActief(alle);
    document.getElementById('zoek-teller').textContent = `1 / ${alle.length}`;
  } else {
    document.getElementById('zoek-teller').textContent = 'geen resultaten';
  }
}

function zoekNavigeer(richting) {
  const alle = doc.querySelectorAll('mark.zoek-mark');
  if (!alle.length) return;
  zoekIndex = (zoekIndex + richting + alle.length) % alle.length;
  markeerActief(alle);
  document.getElementById('zoek-teller').textContent = `${zoekIndex + 1} / ${alle.length}`;
}

function markeerActief(alle) {
  alle.forEach((m, i) => {
    m.style.background = i === zoekIndex ? '#ff9900' : '#fff176';
  });
  alle[zoekIndex].scrollIntoView({ behavior:'smooth', block:'center' });
}

// ── zoek & vervang ───────────────────────────────────────────────────────────

function sluitVervang() {
  document.getElementById('vervang-paneel').classList.remove('zichtbaar');
}

function vervangEen() {
  const zoekT  = document.getElementById('vervang-zoek').value;
  const doorT  = document.getElementById('vervang-door').value;
  if (!zoekT) return;
  const mark = doc.querySelector('mark.zoek-mark');
  if (mark) {
    mark.replaceWith(document.createTextNode(doorT));
    doc.normalize();
    documentGewijzigd();
    zoek(zoekT);
  }
}

function vervangAlles() {
  const zoekT = document.getElementById('vervang-zoek').value;
  const doorT = document.getElementById('vervang-door').value;
  if (!zoekT) return;
  const aantal = doc.querySelectorAll('mark.zoek-mark').length;
  doc.querySelectorAll('mark.zoek-mark').forEach(m => {
    m.replaceWith(document.createTextNode(doorT));
  });
  doc.normalize();
  documentGewijzigd();
  document.getElementById('zoek-teller').textContent = `${aantal} vervangen`;
  document.getElementById('zoek-input').value = '';
}

// ── export ────────────────────────────────────────────────────────────────────

function downloadUrl(url, bestandsnaam) {
  const a = document.createElement('a');
  a.href = url; a.download = bestandsnaam; a.click();
  URL.revokeObjectURL(url);
}

async function exporteerJSON() {
  const data = await api('GET', '/api/export');
  const blob = new Blob([JSON.stringify(data, null, 2)], { type:'application/json' });
  downloadUrl(URL.createObjectURL(blob), `logboek-${datumVandaag()}.json`);
}

async function exporteerHTML() {
  const res  = await fetch('/api/export?formaat=html');
  const blob = await res.blob();
  downloadUrl(URL.createObjectURL(blob), `logboek-${datumVandaag()}.html`);
}

async function exporteerTXT() {
  const res  = await fetch('/api/export?formaat=txt');
  const blob = await res.blob();
  downloadUrl(URL.createObjectURL(blob), `logboek-${datumVandaag()}.txt`);
}

function datumVandaag() { return new Date().toISOString().slice(0,10); }

// ── keyboard shortcuts ────────────────────────────────────────────────────────

document.addEventListener('keydown', e => {
  if (e.ctrlKey || e.metaKey) {
    if (e.key==='b') { e.preventDefault(); fmt('bold'); }
    if (e.key==='i') { e.preventDefault(); fmt('italic'); }
    if (e.key==='u') { e.preventDefault(); fmt('underline'); }
    if (e.key==='s') { e.preventDefault(); bewaar(); }
    if (e.key==='h') {
      e.preventDefault();
      document.getElementById('vervang-paneel').classList.toggle('zichtbaar');
      if (document.getElementById('vervang-paneel').classList.contains('zichtbaar')) {
        document.getElementById('vervang-zoek').focus();
      }
    }
  }
  if (e.key === 'Escape') {
    sluitVervang();
    document.getElementById('zoek-input').value = '';
    zoek('');
  }
});

// ── hulpfuncties ──────────────────────────────────────────────────────────────

async function api(methode, pad, body) {
  const opts = { method: methode, headers: {} };
  if (body) { opts.headers['Content-Type'] = 'application/json'; opts.body = JSON.stringify(body); }
  return (await fetch(pad, opts)).json();
}

function escHtml(s) {
  return s.replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;');
}

function escRegex(s) {
  return s.replace(/[.*+?^${}()|[\]\\]/g, '\\$&');
}

function setStatus(t) {
  document.getElementById('status-balk').textContent = t;
}

function netjesTijd(iso) {
  if (!iso) return '';
  const d = new Date(iso);
  return d.toLocaleTimeString('nl-NL', { hour:'2-digit', minute:'2-digit', second:'2-digit' });
}

// ── afbeelding resizen ────────────────────────────────────────────────────────

let actieveWrap = null;

function selecteerAfbeelding(img) {
  // verwijder eerdere selectie
  deselecteerAlles();

  const wrap = document.createElement('span');
  wrap.className = 'img-resize-wrap';
  wrap.contentEditable = 'false';

  img.parentNode.insertBefore(wrap, img);
  wrap.appendChild(img);

  ['lt','rt','lb','rb'].forEach(pos => {
    const greep = document.createElement('span');
    greep.className = `resize-greep ${pos}`;
    greep.addEventListener('mousedown', e => startResize(e, img, wrap, pos));
    wrap.appendChild(greep);
  });

  img.classList.add('geselecteerd');
  actieveWrap = wrap;
}

function deselecteerAlles() {
  document.querySelectorAll('.img-resize-wrap').forEach(w => {
    const img = w.querySelector('img');
    if (img) {
      img.classList.remove('geselecteerd');
      w.parentNode.insertBefore(img, w);
    }
    w.remove();
  });
  actieveWrap = null;
}

function startResize(e, img, wrap, pos) {
  e.preventDefault();
  e.stopPropagation();

  const startX  = e.clientX;
  const startW  = img.offsetWidth;
  const startH  = img.offsetHeight;
  const ratio   = startH / startW;
  const linksPos = pos === 'lt' || pos === 'lb';

  function onMove(ev) {
    const dx   = linksPos ? startX - ev.clientX : ev.clientX - startX;
    const nieuwW = Math.max(40, startW + dx);
    const nieuwH = Math.round(nieuwW * ratio);
    img.style.width  = nieuwW + 'px';
    img.style.height = nieuwH + 'px';
  }

  function onUp() {
    document.removeEventListener('mousemove', onMove);
    document.removeEventListener('mouseup',   onUp);
    documentGewijzigd();
  }

  document.addEventListener('mousemove', onMove);
  document.addEventListener('mouseup',   onUp);
}

// klik op afbeelding → selecteren; klik ergens anders → deselecteren
doc.addEventListener('click', e => {
  if (e.target.tagName === 'IMG' && doc.contains(e.target)) {
    e.preventDefault();
    selecteerAfbeelding(e.target);
  } else if (!e.target.closest('.img-resize-wrap')) {
    deselecteerAlles();
  }
});

// bij bewaren: verwijder resize-wraps (img behoudt width/height als inline stijl)
const origBewaar = bewaar;
bewaar = async function() {
  deselecteerAlles();
  await origBewaar();
};

init();
</script>
</body>
</html>
"""

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
